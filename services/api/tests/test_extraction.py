import io
import zipfile
import pytest
from atlas.extraction import extract, chunk_blocks, Block


def test_csv_preserves_headers_and_rows():
    blocks, _ = extract(b"employee,limit\nAlice,5000\nBob,3000", "expenses.csv")
    assert "employee: Alice" in blocks[0].text
    assert "limit: 5000" in blocks[0].text
    assert blocks[0].location["start"] == 2
    assert blocks[0].location["end"] == 3


def test_utf16_and_binary_guard():
    blocks, _ = extract("Company policy: 18 days leave".encode("utf-16"), "policy.txt")
    assert "18 days" in blocks[0].text
    with pytest.raises(ValueError, match="binary"):
        extract(b"hello\x00world", "readme.txt")


def test_docx_paragraph_and_table_locations():
    from docx import Document

    document = Document()
    document.add_paragraph("Remote work requires manager approval.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Days"
    table.cell(0, 1).text = "2"
    output = io.BytesIO()
    document.save(output)
    blocks, _ = extract(output.getvalue(), "policy.docx")
    assert blocks[0].location["kind"] == "paragraph"
    assert any("Days | 2" in block.text for block in blocks)


def test_xlsx_maintains_sheet_headers_and_row_groups():
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Expenses"
    sheet.append(["Name", "Amount"])
    for i in range(25):
        sheet.append([f"Employee {i}", i * 10])
    output = io.BytesIO()
    workbook.save(output)
    blocks, note = extract(output.getvalue(), "budget.xlsx")
    assert len(blocks) == 2
    assert all("Columns: Name | Amount" in block.text for block in blocks)
    assert blocks[1].location["sheet"] == "Expenses"
    assert "not recalculated" in note


def test_pptx_keeps_slide_references():
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Security policy"
    slide.placeholders[1].text = "Use MFA for production accounts."
    output = io.BytesIO()
    presentation.save(output)
    blocks, _ = extract(output.getvalue(), "security.pptx")
    assert blocks[0].location["slide"] == 1
    assert "MFA" in blocks[0].text


def test_html_strips_active_content():
    blocks, _ = extract(
        b"<h1>Policy</h1><script>STEAL_SECRET()</script><p>18 leave days</p>",
        "policy.html",
    )
    assert "18 leave days" in blocks[0].text
    assert "STEAL_SECRET" not in blocks[0].text


def test_chunking_is_bounded_and_retains_lines():
    text = "\n".join(f"Line {i}: " + "x" * 60 for i in range(100))
    chunks = chunk_blocks(
        [
            Block(
                text,
                {"kind": "lines", "start": 10, "end": 109, "label": "Lines 10–109"},
            )
        ]
    )
    assert all(len(chunk.text) <= 2200 for chunk in chunks)
    assert chunks[-1].location["end"] == 109
    assert chunks[1].location["start"] > 10
    assert chunks[-1].text.endswith("x" * 60)


def test_empty_and_unsupported_files_fail_explicitly():
    with pytest.raises(ValueError, match="No readable text"):
        extract(b"   ", "empty.txt")
    with pytest.raises(ValueError, match="Unsupported"):
        extract(b"MZ", "run.exe")


def test_pdf_extraction_keeps_page_citations():
    from atlas.examples import sample_pdf

    blocks, note = extract(sample_pdf(), "travel.pdf")
    assert "14 days" in blocks[0].text
    assert blocks[0].location["page"] == 1
    assert "All 1 PDF pages read" in note


def test_zip_keeps_member_locations_without_extracting_to_disk():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("engineering/setup.md", "Install dependencies using npm ci.")
        archive.writestr("policy.txt", "Employees get 18 leave days.")
        archive.writestr("binary.exe", b"MZ")
    blocks, note = extract(buffer.getvalue(), "company.zip")
    assert blocks[0].location["archive_member"] == "engineering/setup.md"
    assert "1 unsupported" in note
    assert len(blocks) == 2


def test_zip_rejects_traversal_and_nested_archives():
    for name in ("../secret.txt", "/absolute.txt", "nested.zip"):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            archive.writestr(name, b"unsafe")
        with pytest.raises(ValueError):
            extract(buffer.getvalue(), "company.zip")
