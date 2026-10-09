"""Bounded parsers. Location metadata remains attached to every extracted block."""

import csv
import io
import json
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from charset_normalizer import from_bytes

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".log",
    ".py",
    ".js",
    ".jsx",
    ".ts",
    ".tsx",
    ".java",
    ".c",
    ".h",
    ".cpp",
    ".cs",
    ".go",
    ".rs",
    ".rb",
    ".php",
    ".swift",
    ".kt",
    ".scala",
    ".r",
    ".sh",
    ".ps1",
    ".sql",
    ".json",
    ".yaml",
    ".yml",
    ".toml",
    ".ini",
    ".env",
    ".cfg",
    ".xml",
    ".css",
    ".scss",
    ".html",
    ".htm",
    ".ipynb",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {
    ".pdf",
    ".docx",
    ".xlsx",
    ".pptx",
    ".csv",
    ".tsv",
    ".png",
    ".jpg",
    ".jpeg",
    ".tif",
    ".tiff",
    ".webp",
    ".zip",
}
MAX_TEXT = 8_000_000


@dataclass
class Block:
    text: str
    location: dict


def decode(data: bytes) -> str:
    if b"\x00" in data[:4096] and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise ValueError(
            "This file contains binary data. Upload a supported document or text export."
        )
    detected = from_bytes(data).best()
    if detected is None:
        raise ValueError(
            "The text encoding could not be read. Export the file as UTF-8."
        )
    return str(detected)


def validate_office(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(e.file_size for e in entries) > 100_000_000:
            raise ValueError(
                "This Office document expands beyond the safe processing limit."
            )
        if any(e.flag_bits & 1 for e in entries):
            raise ValueError(
                "Encrypted documents are not supported. Upload a decrypted copy."
            )


def ocr_image(data: bytes) -> str:
    from PIL import Image
    import pytesseract

    if not shutil.which("tesseract"):
        raise ValueError(
            "OCR is unavailable on this server. Upload a text-based export."
        )
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > 25_000_000:
            raise ValueError("Image resolution exceeds the safe OCR limit.")
        return pytesseract.image_to_string(image, timeout=30)


def extract(data: bytes, name: str) -> tuple[list[Block], str | None]:
    suffix = Path(name).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported format. Upload PDF, Office documents, text, code, tables, or images."
        )
    blocks: list[Block] = []
    note = None
    if suffix == ".zip":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
            if (
                len(entries) > 100
                or sum(entry.file_size for entry in entries) > 75_000_000
            ):
                raise ValueError(
                    "Archives are limited to 100 files and 75 MB of expanded content."
                )
            skipped = 0
            for entry in entries:
                path = PurePosixPath(entry.filename.replace("\\", "/"))
                if (
                    path.is_absolute()
                    or ".." in path.parts
                    or ":" in entry.filename
                    or ((entry.external_attr >> 16) & 0o170000) == 0o120000
                ):
                    raise ValueError(
                        "Archive contains an unsafe path or symbolic link."
                    )
                if entry.flag_bits & 1:
                    raise ValueError("Encrypted archives are not supported.")
                if (
                    entry.file_size > 25_000_000
                    or entry.file_size / max(entry.compress_size, 1) > 250
                ):
                    raise ValueError(
                        "An archive member exceeds the safe size or expansion limit."
                    )
                if path.suffix.lower() == ".zip":
                    raise ValueError(
                        "Nested ZIP archives are not supported. Upload their contents separately."
                    )
                if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
                    skipped += 1
                    continue
                extracted, _ = extract(archive.read(entry), str(path))
                for block in extracted:
                    block.location = {
                        **block.location,
                        "archive_member": str(path),
                        "label": str(path) + " · " + block.location["label"],
                    }
                    blocks.append(block)
            note = "Archive members retain their individual source locations."
            if skipped:
                note += f" {skipped} unsupported files were skipped."
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError(
                "Encrypted PDFs are not supported. Upload a decrypted copy."
            )
        if len(reader.pages) > 200:
            raise ValueError("PDFs are limited to 200 pages per upload.")
        scanned = []
        for index, page in enumerate(reader.pages):
            text = page.extract_text(extraction_mode="layout") or ""
            if len(text.strip()) < 30:
                scanned.append(index)
            blocks.append(
                Block(
                    text,
                    {"kind": "page", "page": index + 1, "label": f"Page {index + 1}"},
                )
            )
        if scanned:
            if len(scanned) > 40:
                raise ValueError(
                    "Scanned PDFs are limited to 40 OCR pages. Split the document first."
                )
            if not shutil.which("pdftoppm") or not shutil.which("tesseract"):
                raise ValueError(
                    "This PDF requires OCR, which is unavailable on this server."
                )
            with tempfile.TemporaryDirectory() as directory:
                source = Path(directory) / "source.pdf"
                source.write_bytes(data)
                for index in scanned:
                    dest = Path(directory) / f"page-{index}"
                    subprocess.run(
                        [
                            "pdftoppm",
                            "-f",
                            str(index + 1),
                            "-l",
                            str(index + 1),
                            "-scale-to",
                            "2000",
                            "-singlefile",
                            "-png",
                            str(source),
                            str(dest),
                        ],
                        check=True,
                        timeout=40,
                        capture_output=True,
                    )
                    blocks[index] = Block(
                        ocr_image(dest.with_suffix(".png").read_bytes()),
                        {
                            "kind": "page",
                            "page": index + 1,
                            "ocr": True,
                            "label": f"Page {index + 1} · OCR",
                        },
                    )
            note = "Some pages were read using OCR; verify important numbers against the original."
    elif suffix in {".docx", ".xlsx", ".pptx"}:
        validate_office(data)
        if suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(data))
            for index, paragraph in enumerate(document.paragraphs):
                if paragraph.text.strip():
                    blocks.append(
                        Block(
                            paragraph.text,
                            {
                                "kind": "paragraph",
                                "paragraph": index + 1,
                                "label": f"Paragraph {index + 1}",
                            },
                        )
                    )
            for index, table in enumerate(document.tables):
                text = "\n".join(
                    " | ".join(c.text for c in row.cells) for row in table.rows
                )
                blocks.append(
                    Block(text, {"kind": "table", "label": f"Table {index + 1}"})
                )
        elif suffix == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            total_rows = 0
            for sheet in workbook:
                group = []
                header = ""
                start = 1
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    total_rows += 1
                    if total_rows > 100000:
                        raise ValueError("Spreadsheets are limited to 100,000 rows.")
                    line = " | ".join(
                        "" if value is None else str(value) for value in row
                    )
                    if row_index == 1:
                        header = line
                    group.append(f"Row {row_index}: {line}")
                    if len(group) >= 20:
                        blocks.append(
                            Block(
                                f"Sheet: {sheet.title}\nColumns: {header}\n"
                                + "\n".join(group),
                                {
                                    "kind": "sheet",
                                    "sheet": sheet.title,
                                    "start": start,
                                    "end": row_index,
                                    "label": f"{sheet.title} · rows {start}–{row_index}",
                                },
                            )
                        )
                        group, start = [], row_index + 1
                if group:
                    blocks.append(
                        Block(
                            f"Sheet: {sheet.title}\nColumns: {header}\n"
                            + "\n".join(group),
                            {
                                "kind": "sheet",
                                "sheet": sheet.title,
                                "start": start,
                                "end": start + len(group) - 1,
                                "label": f"{sheet.title} · rows {start}–{start + len(group) - 1}",
                            },
                        )
                    )
            workbook.close()
            note = "Spreadsheet values reflect the saved workbook. Formulas are not recalculated."
        else:
            from pptx import Presentation

            for index, slide in enumerate(Presentation(io.BytesIO(data)).slides):
                text = []
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text.append(shape.text)
                    if shape.has_table:
                        text.extend(
                            " | ".join(cell.text for cell in row.cells)
                            for row in shape.table.rows
                        )
                if slide.has_notes_slide:
                    text.append(slide.notes_slide.notes_text_frame.text)
                blocks.append(
                    Block(
                        "\n".join(text),
                        {
                            "kind": "slide",
                            "slide": index + 1,
                            "label": f"Slide {index + 1}",
                        },
                    )
                )
    elif suffix in {".csv", ".tsv"}:
        text = decode(data)
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
        rows = csv.reader(io.StringIO(text), dialect=dialect)
        header = next(rows, [])
        group, start = [], 2
        for index, row in enumerate(rows, 2):
            if index > 100000:
                raise ValueError("Tables are limited to 100,000 rows.")
            group.append(
                f"Row {index}: "
                + " | ".join(
                    f"{header[j] if j < len(header) else j + 1}: {value}"
                    for j, value in enumerate(row)
                )
            )
            if len(group) >= 20:
                blocks.append(
                    Block(
                        "Columns: " + " | ".join(header) + "\n" + "\n".join(group),
                        {
                            "kind": "rows",
                            "start": start,
                            "end": index,
                            "label": f"Rows {start}–{index}",
                        },
                    )
                )
                group, start = [], index + 1
        if group:
            blocks.append(
                Block(
                    "Columns: " + " | ".join(header) + "\n" + "\n".join(group),
                    {
                        "kind": "rows",
                        "start": start,
                        "end": start + len(group) - 1,
                        "label": f"Rows {start}–{start + len(group) - 1}",
                    },
                )
            )
        if not blocks and header:
            blocks.append(
                Block(" | ".join(header), {"kind": "rows", "label": "Header row"})
            )
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}:
        blocks.append(
            Block(
                ocr_image(data), {"kind": "image", "ocr": True, "label": "Image · OCR"}
            )
        )
        note = "Image text was read using OCR; verify important details against the original."
    else:
        text = decode(data)
        if suffix in {".html", ".htm", ".xml"}:
            from bs4 import BeautifulSoup

            soup = BeautifulSoup(text, "html.parser")
            for element in soup(["script", "style", "iframe", "object"]):
                element.decompose()
            text = soup.get_text("\n", strip=True)
        if suffix == ".ipynb":
            parsed = json.loads(text)
            text = "\n\n".join(
                "".join(cell.get("source", [])) for cell in parsed.get("cells", [])
            )
        lines = text.splitlines()
        for start in range(0, len(lines), 45):
            end = min(start + 45, len(lines))
            blocks.append(
                Block(
                    "\n".join(lines[start:end]),
                    {
                        "kind": "lines",
                        "start": start + 1,
                        "end": end,
                        "label": f"Lines {start + 1}–{end}",
                    },
                )
            )
    blocks = [block for block in blocks if block.text.strip()]
    if not blocks:
        raise ValueError(
            "No readable text was found. Upload a clearer image or a text-based export."
        )
    if sum(len(block.text) for block in blocks) > MAX_TEXT:
        raise ValueError(
            "The extracted content exceeds the processing limit. Split the file."
        )
    return blocks, note


def chunk_blocks(
    blocks: list[Block], max_chars: int = 2200, overlap: int = 200
) -> list[Block]:
    chunks = []
    for block in blocks:
        text = block.text.rstrip()
        start = 0
        while start < len(text):
            end = min(start + max_chars, len(text))
            if end < len(text):
                boundary = text.rfind("\n", start + max_chars // 2, end)
                if boundary > start:
                    end = boundary
            location = dict(block.location)
            # Character slices inside long source blocks retain truthful line ranges.
            if location.get("kind") == "lines":
                location["start"] = block.location["start"] + text[:start].count("\n")
                location["end"] = location["start"] + text[start:end].count("\n")
                prefix = (
                    (location["archive_member"] + " · ")
                    if location.get("archive_member")
                    else ""
                )
                location["label"] = (
                    prefix + f"Lines {location['start']}–{location['end']}"
                )
            chunks.append(Block(text[start:end], location))
            if end == len(text):
                break
            start = max(start + 1, end - overlap)
    return chunks
