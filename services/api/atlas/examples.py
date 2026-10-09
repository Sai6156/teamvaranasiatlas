"""Small fictional documents for an optional workspace quickstart."""

import io


def sample_pdf() -> bytes:
    lines = [
        "FICTIONAL EXAMPLE - Atlas Travel Policy",
        "Business travel requires manager approval before booking.",
        "Expenses are reimbursed up to INR 5000 per day with receipts.",
        "Submit travel expense claims within 14 days of the trip.",
        "This is synthetic sample data, not a real company policy.",
    ]
    instructions = (
        "BT /F1 13 Tf 50 780 Td "
        + " ".join(f"({line}) Tj 0 -28 Td" for line in lines)
        + " ET"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(instructions)} >>\nstream\n{instructions}\nendstream".encode(),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return bytes(output)


def starter_files():
    from docx import Document

    document = Document()
    document.add_heading("Fictional Example: Onboarding", 0)
    document.add_paragraph(
        "On your first day, collect your laptop from IT, activate multifactor authentication, and meet your buddy. Complete security training within five working days."
    )
    buffer = io.BytesIO()
    document.save(buffer)
    return [
        (
            "Example - People handbook.md",
            b"# Fictional example: People handbook\nThis is synthetic sample data.\nEmployees receive 18 days of annual leave each calendar year. Request leave through the HR portal at least 5 working days in advance. Your manager approves the request.\nHybrid working is available for two days each week with manager approval.\n",
            "People & policies",
        ),
        ("Example - Travel policy.pdf", sample_pdf(), "Operations"),
        ("Example - Onboarding.docx", buffer.getvalue(), "People & policies"),
        (
            "Example - Expense limits.csv",
            b"category,daily_limit_inr,approval\nmeals,800,manager\nhotel,4000,manager\ntaxi,1200,manager\n",
            "Operations",
        ),
        (
            "Example - Engineering guide.py",
            b"# Fictional example: engineering setup\n# Use Python 3.12. Install dependencies with pip install -r requirements.txt.\n# Run tests with pytest. Never commit API keys or .env files.\n# Production access requires multifactor authentication and team lead approval.\ndef health():\n    return {'status': 'ok'}\n",
            "Engineering",
        ),
    ]
