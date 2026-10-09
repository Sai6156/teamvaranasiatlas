"""Live synthetic acceptance test. Secrets and generated credentials stay local."""
import asyncio
import io
import json
import secrets
import sys
import time
from pathlib import Path
import httpx

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "services" / "api"))
from atlas.config import Settings

config = Settings(_env_file=(ROOT / "services/api/.env", ROOT / "secrets.env"))
URL = config.supabase_url
API = sys.argv[1] if len(sys.argv) > 1 else "http://127.0.0.1:8000"
ADMIN = {"apikey": config.supabase_service_role_key}
if not config.supabase_service_role_key.startswith("sb_secret_"):
    ADMIN["Authorization"] = "Bearer " + config.supabase_service_role_key
PUBLIC = {"apikey": config.supabase_publishable_key}


async def main():
    suffix = secrets.token_hex(4)
    accounts = []
    async with httpx.AsyncClient(timeout=90) as client:
        for label in ("admin", "outsider", "employee"):
            email = f"atlas-qa-{label}-{suffix}@example.com"
            password = secrets.token_urlsafe(20)
            response = await client.post(URL+"/auth/v1/admin/users", headers=ADMIN, json={"email": email, "password": password, "email_confirm": True, "user_metadata": {"full_name": "Atlas QA " + label}})
            response.raise_for_status()
            user = response.json()
            login = await client.post(URL+"/auth/v1/token?grant_type=password", headers=PUBLIC, json={"email": email, "password": password})
            login.raise_for_status()
            accounts.append({"id": user["id"], "email": email, "password": password, "token": login.json()["access_token"]})
        a, b, employee = accounts
        auth_a, auth_b, auth_e = [{"Authorization": "Bearer " + u["token"]} for u in accounts]
        result = await client.post(API+"/workspaces", headers=auth_a, json={"name": "Atlas QA Company " + suffix, "full_name": "Atlas QA Admin"})
        result.raise_for_status()
        org = result.json()["id"]
        outsider = await client.post(API+"/workspaces", headers=auth_b, json={"name": "Isolated QA Company " + suffix})
        outsider.raise_for_status()
        other_org = outsider.json()["id"]
        assert (await client.get(API+f"/workspaces/{org}/documents", headers=auth_b)).status_code == 403
        assert (await client.post(API+f"/workspaces/{org}/documents", headers=auth_b, files={"file": ("stolen.txt", b"No access")})).status_code == 403
        invitation = await client.post(API+f"/workspaces/{org}/invitations", headers=auth_a, json={"email": employee["email"], "role": "employee"})
        assert invitation.status_code == 201, invitation.json()
        token = invitation.json()["url"].split("invite=")[1]
        assert (await client.post(API+"/invitations/accept", headers=auth_b, json={"token": token})).status_code == 400
        accepted = await client.post(API+"/invitations/accept", headers=auth_e, json={"token": token})
        accepted.raise_for_status()
        assert (await client.post(API+f"/workspaces/{org}/documents", headers=auth_e, files={"file": ("forbidden.txt", b"Employee cannot upload")})).status_code == 403
        from docx import Document
        document = Document()
        document.add_heading("Travel policy", 0)
        document.add_paragraph("Employees can claim business travel expenses up to INR 5000 per day with receipts. Claims must be submitted within 14 days. Manager approval is required before booking.")
        office = io.BytesIO()
        document.save(office)
        files = [
            ("people-handbook.md", b"# People handbook\nEmployees receive 18 days of annual leave per calendar year. Request leave through the HR portal at least 5 working days in advance. Your manager approves the request. On your first day, collect your laptop from IT and enable multifactor authentication.\n", "People & policies"),
            ("travel-policy.docx", office.getvalue(), "Operations"),
            ("expense-limits.csv", b"category,limit_inr,approval\nmeals,800,manager\nhotel,4000,manager\ntaxi,1200,manager\n", "Operations"),
            ("engineering-guide.py", b"# Development setup\n# Use Python 3.12. Install requirements with pip install -r requirements.txt.\n# Run tests with pytest. Never commit .env or API keys.\ndef health():\n    return {'status': 'ok'}\n", "Engineering"),
        ]
        doc_ids = []
        for name, content, collection in files:
            response = await client.post(API+f"/workspaces/{org}/documents", headers=auth_a, data={"collection": collection}, files={"file": (name, content, "application/octet-stream")})
            response.raise_for_status()
            doc_ids.append(response.json()["id"])
        duplicate = await client.post(API+f"/workspaces/{org}/documents", headers=auth_a, files={"file": (files[0][0], files[0][1])})
        assert duplicate.status_code == 409
        deadline = time.monotonic() + 180
        while time.monotonic() < deadline:
            response = await client.get(API+f"/workspaces/{org}/documents", headers=auth_a)
            response.raise_for_status()
            documents = response.json()
            failures = [d for d in documents if d["status"] == "failed"]
            if failures:
                raise RuntimeError("Indexing failed: " + "; ".join(d.get("error_message", "") for d in failures))
            if len(documents) == 4 and all(d["status"] == "ready" for d in documents):
                break
            await asyncio.sleep(3)
        else:
            raise RuntimeError("Indexing did not complete within the acceptance-test window")
        assert (await client.get(API+f"/documents/{doc_ids[0]}/source", headers=auth_b)).status_code == 404
        assert (await client.delete(API+f"/documents/{doc_ids[0]}", headers=auth_e)).status_code == 400
        signed = await client.get(API+f"/documents/{doc_ids[0]}/source", headers=auth_e)
        signed.raise_for_status()
        assert (await client.get(signed.json()["url"])).status_code == 200
        response = await client.post(API+f"/workspaces/{org}/chat", headers=auth_e, json={"question": "How many days of annual leave do employees get, and within how many days must travel expense claims be submitted?"})
        response.raise_for_status()
        events = [json.loads(line[5:].strip()) for line in response.text.splitlines() if line.startswith("data:")]
        errors = [event for event in events if event["type"] == "error"]
        assert not errors, errors
        final = next(event for event in events if event["type"] == "done")
        answer = "".join(event["text"] for event in events if event["type"] == "token")
        assert "18" in answer and "14" in answer, "Expected leave and expense facts were not present"
        assert len({c["document_id"] for c in final["citations"]}) >= 2, "Expected cross-document citations"
        conversation = next(event["id"] for event in events if event["type"] == "conversation")
        assert (await client.get(API+f"/conversations/{conversation}/messages", headers=auth_b)).status_code == 404
        assert (await client.get(API+f"/conversations/{conversation}/messages", headers=auth_a)).status_code == 404
        # Keep synthetic QA data for UI verification; a separate cleanup removes it.
        runtime = ROOT / ".runtime"
        runtime.mkdir(exist_ok=True)
        (runtime / "qa-session.json").write_text(json.dumps({"accounts": accounts, "organizations": [org, other_org], "documents": doc_ids}), encoding="utf-8")
        report = {"passed": True, "documents_indexed": len(documents), "cross_document_sources": len(final["citations"]), "model": final["model"], "question_duration_ms": final["duration_ms"], "isolation_checks": 8}
        (runtime / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(json.dumps(report))


asyncio.run(main())
